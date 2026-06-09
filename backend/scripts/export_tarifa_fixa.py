"""Exporta planilha com tarifa fixa do ML pra TODOS os anúncios ativos do perfil.

Regra observada do ML BR (mai/2026):
  - free_shipping=true OU preço >= R$ 79: tarifa fixa = 0
  - preço < R$ 12,50: tarifa fixa = 50% do preço
  - listing_type_id="gold_pro" (Premium): R$ 8,55
  - listing_type_id="gold_special" (Clássico) e outras: R$ 6,75

Inclui também comissão (% e R$), custo do xlsx do perfil, "você recebe" e
margem líquida estimada — pra reconciliação com o painel do ML.

Uso:
  uv run --project backend python backend/scripts/export_tarifa_fixa.py [profile_slug]

Default: namore (loja do exemplo do user). Aceita slug pra rodar pra outras lojas.
"""
from __future__ import annotations

import asyncio
import sys
from pathlib import Path
from typing import Any
from uuid import UUID

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from liraz_tools.infrastructure.db.database import session_scope
from liraz_tools.infrastructure.ml.client import MLClient
from liraz_tools.infrastructure.ml.items_fetcher import _extrair_sku
from liraz_tools.infrastructure.ml.promotions_lookup import (
    _listar_todos_items_ativos_do_vendedor,
)
from liraz_tools.infrastructure.pricing.costs_loader import (
    buscar_custo,
    carregar_custos,
)
from liraz_tools.infrastructure.repositories.cost_overrides_repository import (
    CostOverridesRepository,
)
from liraz_tools.infrastructure.repositories.per_profile_credentials_repository import (
    PerProfileCredentialsRepository,
)
from liraz_tools.infrastructure.repositories.profile_repository import (
    SQLAlchemyProfileRepository,
)

SLUG_PARA_ID = {
    "namore": UUID("722d19aa-983b-4a18-8cb9-70484520fcda"),
    "liraz": UUID("7de27552-bd0e-4f5a-a03a-d5206177f325"),
    "toque-rico": UUID("94c5477b-899e-4019-b651-4be75213e6fd"),
}

DEGRAU = 79.0
TETO_CLASSICO = 6.75
TETO_PREMIUM = 8.55
LIMITE_BAIXO = 12.50
PCT_BAIXO = 0.50

ATRIBUTOS_MULTIGET = ",".join([
    "id", "title", "price", "original_price",
    "available_quantity", "status",
    "listing_type_id", "category_id", "shipping",
    "seller_custom_field", "attributes", "variations",
])
BATCH = 20  # multiget /items aceita até 20 ids


def calcular_tarifa_fixa(
    *, preco: float, free_shipping: bool, listing_type_id: str,
) -> tuple[float, str]:
    """Aplica a regra observada. Retorna (valor, fonte/motivo)."""
    if preco <= 0:
        return 0.0, "preco_invalido"
    if free_shipping or preco >= DEGRAU:
        return 0.0, "isento (frete gratis ou >=R$79)"
    if preco < LIMITE_BAIXO:
        return round(preco * PCT_BAIXO, 2), f"50% do preco (< R${LIMITE_BAIXO})"
    if listing_type_id == "gold_pro":
        return TETO_PREMIUM, "Premium (gold_pro)"
    return TETO_CLASSICO, f"Classico/{listing_type_id}"


async def _fetch_listing_prices(
    ml: MLClient, *, preco: float, category_id: str, listing_type_id: str,
    logistic_type: str, shipping_mode: str,
) -> dict[str, Any] | None:
    """Lê fees teóricos do ML (comissão %) — 1 call por item."""
    try:
        resp = await ml.get(
            "/sites/MLB/listing_prices",
            params={
                "price": preco,
                "category_id": category_id,
                "listing_type_id": listing_type_id,
                "currency_id": "BRL",
                "logistic_type": logistic_type or "cross_docking",
                "shipping_mode": shipping_mode or "me2",
            },
            use_cache=True,
        )
    except Exception:
        return None
    if isinstance(resp, list):
        for e in resp:
            if isinstance(e, dict) and e.get("listing_type_id") == listing_type_id:
                return e
        return resp[0] if resp else None
    return resp if isinstance(resp, dict) else None


async def main() -> None:
    slug = sys.argv[1] if len(sys.argv) > 1 else "namore"
    if slug not in SLUG_PARA_ID:
        print(f"Slug invalido: {slug}. Opcoes: {list(SLUG_PARA_ID.keys())}")
        sys.exit(1)
    profile_id = SLUG_PARA_ID[slug]

    async with session_scope() as session:
        profile_repo = SQLAlchemyProfileRepository(session)
        profile = await profile_repo.get_by_id(profile_id)
        assert profile.ml_user_id is not None

        creds_repo = PerProfileCredentialsRepository()
        creds = creds_repo.get_app_credentials(profile.slug)
        tokens = creds_repo.get_tokens(profile.slug, profile.ml_user_id)
        overrides_repo = CostOverridesRepository()
        aliquota = profile.config.aliquota_imposto

        # Carrega custos (xlsx do perfil + overrides manuais) — mesma lógica
        # que o sugestao_use_case usa.
        custos_xlsx: dict[str, float] = {}
        if profile.config.custos_xlsx_path:
            try:
                custos_xlsx = carregar_custos(profile.config.custos_xlsx_path)
            except Exception as e:
                print(f"  AVISO: falhou ao carregar xlsx: {e}")
        try:
            overrides = overrides_repo.load_all(profile.slug)
        except Exception:
            overrides = {}

        async with MLClient(creds, tokens) as ml:
            print(f"[{profile.name}] Listando anuncios ativos...")
            ativos = await _listar_todos_items_ativos_do_vendedor(
                ml, profile.ml_user_id,
            )
            print(f"  {len(ativos)} anuncios ativos.")

            # Multiget pra info básica
            print(f"[{profile.name}] Multiget de info basica em batches de {BATCH}...")
            batches = [ativos[i:i + BATCH] for i in range(0, len(ativos), BATCH)]

            async def _multiget(batch: list[str]) -> list[dict[str, Any]]:
                try:
                    resp = await ml.get(
                        "/items",
                        params={
                            "ids": ",".join(batch),
                            "attributes": ATRIBUTOS_MULTIGET,
                        },
                    )
                except Exception:
                    return []
                return resp if isinstance(resp, list) else []

            mg_results = await asyncio.gather(*[_multiget(b) for b in batches])
            items_raw: list[dict[str, Any]] = []
            for batch in mg_results:
                for entry in batch:
                    if (
                        isinstance(entry, dict) and entry.get("code") == 200
                        and isinstance(entry.get("body"), dict)
                    ):
                        items_raw.append(entry["body"])
            print(f"  {len(items_raw)} items hidratados.")

            # Listing prices (comissão %) — paralelo, com cache no MLClient
            print(f"[{profile.name}] Consultando listing_prices (comissao)...")

            async def _lp_um(it: dict[str, Any]) -> tuple[str, dict[str, Any] | None]:
                shipping = it.get("shipping") or {}
                lp = await _fetch_listing_prices(
                    ml, preco=float(it.get("price") or 0),
                    category_id=str(it.get("category_id") or ""),
                    listing_type_id=str(it.get("listing_type_id") or "gold_special"),
                    logistic_type=str(shipping.get("logistic_type") or ""),
                    shipping_mode=str(shipping.get("mode") or ""),
                )
                return it.get("id", ""), lp

            lp_pairs = await asyncio.gather(*[_lp_um(it) for it in items_raw])
            mapa_lp = {iid: lp for iid, lp in lp_pairs if iid}
            print(f"  {sum(1 for v in mapa_lp.values() if v)} listing_prices lidos.")

    # Monta linhas da planilha
    print("\nMontando planilha...")
    linhas: list[dict[str, Any]] = []
    for it in items_raw:
        item_id = str(it.get("id") or "")
        sku = _extrair_sku(it)
        titulo = str(it.get("title") or "")
        preco = float(it.get("price") or 0)
        status = str(it.get("status") or "")
        modalidade_id = str(it.get("listing_type_id") or "")
        modalidade_nome = {
            "gold_special": "Classico",
            "gold_pro": "Premium",
            "gold": "Ouro",
            "silver": "Prata",
            "bronze": "Bronze",
            "free": "Gratuito",
        }.get(modalidade_id, modalidade_id)
        shipping = it.get("shipping") or {}
        free = bool(shipping.get("free_shipping"))

        tarifa_fixa, tarifa_fonte = calcular_tarifa_fixa(
            preco=preco, free_shipping=free,
            listing_type_id=modalidade_id,
        )

        lp = mapa_lp.get(item_id)
        comissao_pct = 0.0
        comissao_valor = 0.0
        if isinstance(lp, dict):
            sale_fee = lp.get("sale_fee_details") or {}
            comissao_pct = float(sale_fee.get("percentage_fee") or 0)
            comissao_valor = float(lp.get("sale_fee_amount") or 0)
        if comissao_valor == 0 and comissao_pct > 0:
            comissao_valor = round(preco * comissao_pct / 100, 2)

        custo_unit, custo_fonte = buscar_custo(sku, item_id, custos_xlsx, overrides)
        voce_recebe = round(preco - comissao_valor - tarifa_fixa, 2)
        imposto_valor = round(preco * aliquota, 2)
        if custo_unit is not None and preco > 0:
            margem_pct = (
                preco - comissao_valor - tarifa_fixa - imposto_valor - custo_unit
            ) / preco
        else:
            margem_pct = None

        linhas.append({
            "MLB": item_id, "SKU": sku, "Titulo": titulo, "Status": status,
            "Modalidade": modalidade_nome, "Modalidade ID": modalidade_id,
            "Free shipping": free, "Preco": preco,
            "Comissao %": comissao_pct, "Comissao R$": comissao_valor,
            "Tarifa fixa R$": tarifa_fixa, "Tarifa fonte": tarifa_fonte,
            "Aliquota imposto %": round(aliquota * 100, 2),
            "Imposto R$": imposto_valor,
            "Voce recebe (preco - comissao - tarifa_fixa)": voce_recebe,
            "Custo unit": custo_unit,
            "Custo fonte": custo_fonte,
            "Margem liquida %": (
                round(margem_pct * 100, 2) if margem_pct is not None else None
            ),
        })

    # Exporta .xlsx
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "Tarifa fixa"

    if not linhas:
        print("Sem itens. Nada a exportar.")
        return

    headers = list(linhas[0].keys())
    ws.append(headers)
    for col_idx, _ in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col_idx)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="305496")
        cell.alignment = Alignment(horizontal="center")

    for linha in linhas:
        ws.append([linha[h] for h in headers])

    # Larguras razoáveis
    larguras = {
        "MLB": 14, "SKU": 14, "Titulo": 50, "Status": 10,
        "Modalidade": 12, "Modalidade ID": 14, "Free shipping": 10,
        "Preco": 10, "Comissao %": 11, "Comissao R$": 12,
        "Tarifa fixa R$": 14, "Tarifa fonte": 28,
        "Aliquota imposto %": 14, "Imposto R$": 11,
        "Voce recebe (preco - comissao - tarifa_fixa)": 38,
        "Custo unit": 14, "Custo fonte": 14, "Margem liquida %": 14,
    }
    for col_idx, h in enumerate(headers, start=1):
        ws.column_dimensions[get_column_letter(col_idx)].width = larguras.get(h, 14)

    ws.freeze_panes = "A2"

    out = Path(__file__).resolve().parents[2] / f"tarifa_fixa_{slug}.xlsx"
    wb.save(out)
    print(f"\nPlanilha exportada: {out}")
    print(f"Tamanho: {out.stat().st_size} bytes  |  Linhas: {len(linhas)}")

    # Resumo no console
    total = len(linhas)
    cobrando = sum(1 for x in linhas if x["Tarifa fixa R$"] > 0)
    isentos = total - cobrando
    com_custo = sum(1 for x in linhas if x["Custo unit"] is not None)
    print(f"\nResumo:")
    print(f"  Total de anuncios: {total}")
    print(f"  Cobram tarifa fixa: {cobrando}  ({cobrando / total * 100:.1f}%)")
    print(f"  Isentos (free shipping ou >=R$79): {isentos}")
    print(f"  Com custo unitario cadastrado: {com_custo}")


if __name__ == "__main__":
    asyncio.run(main())
