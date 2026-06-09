"""Renderizador do XLSX consolidado com taxas/custos/lucro.

Portado de `tools/relatorio.py` (funções `_criar_workbook` e cabeçalhos).
Layout 100% idêntico ao MCP — mesmas cores, mesmas colunas, mesmas fórmulas.

Colunas:
  A: SKU
  B: ID do Produto (MLB)
  C: Valor de venda
  D: Tarifa Clássico
  E: Tarifa Premium
  F: % da Tarifa
  G: Tarifa Fixa
  H: Frete
  I: Valor Final (fórmula: C-D-E-G-H)
  J: Custo do Produto
  K: Lucro Bruto (fórmula: I-J)
  L: Imposto % (referência à R1)
  M: Imposto R$ (fórmula: C*L)
  N: Lucro Líquido (fórmula: K-M)
  O: % líquida (fórmula: N/C)

Célula R1 contém a alíquota configurável — editar manualmente recalcula tudo.
"""
from __future__ import annotations

from io import BytesIO
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from liraz_tools.infrastructure.pricing.costs_loader import buscar_custo


def render_fee_report_xlsx(
    resultados: list[dict[str, Any]],
    custos: dict[str, float],
    aliquota: float,
    custos_manuais: dict[str, float] | None = None,
) -> tuple[bytes, dict[str, Any]]:
    """Gera o XLSX em memória com layout do relatório consolidado.

    Returns:
        Tuple (xlsx_bytes, diagnostico). Diagnóstico contém contagem
        por tipo de match e listas de SKUs com fallback ou sem custo.
    """
    wb = Workbook()
    ws = wb.active
    if ws is None:  # pragma: no cover - openpyxl sempre cria active
        raise RuntimeError("Workbook sem worksheet ativa")
    ws.title = "Taxas e Lucro"

    # === Estilos (idênticos ao MCP) ===
    font_header = Font(name="Arial", size=11, bold=True)
    font_data = Font(name="Arial", size=10)

    fill_tarifa = PatternFill("solid", start_color="FBE5D6")
    fill_pct = PatternFill("solid", start_color="FFE699")
    fill_fixa = PatternFill("solid", start_color="E2EFDA")
    fill_frete = PatternFill("solid", start_color="DEEBF7")
    fill_final = PatternFill("solid", start_color="FFD966")
    fill_custo = PatternFill("solid", start_color="FFF2CC")
    fill_lucro_bruto = PatternFill("solid", start_color="FFC000")
    fill_imposto = PatternFill("solid", start_color="F4B084")
    fill_lucro_liq = PatternFill("solid", start_color="00B050")
    fill_pct_liq = PatternFill("solid", start_color="92D050")
    fill_header_yellow = PatternFill("solid", start_color="FFFF00")
    fill_warning = PatternFill("solid", start_color="F8CBAD")
    fill_config = PatternFill("solid", start_color="D9E1F2")

    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    left = Alignment(horizontal="left", vertical="center")
    right = Alignment(horizontal="right", vertical="center")

    thin = Side(border_style="thin", color="000000")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    # === Linha 1: cabeçalhos principais (com merges) ===
    headers_l1 = [
        ("D1:E1", "Tarifa", fill_tarifa),
        ("F1", "% da Tarifa", fill_pct),
        ("G1", "Tarifa Fixa", fill_fixa),
        ("H1", "Frete", fill_frete),
        ("I1", "Valor Final", fill_final),
        ("J1", "Custo do Produto", fill_custo),
        ("K1", "Lucro bruto", fill_lucro_bruto),
        ("L1:M1", "Imposto", fill_imposto),
        ("N1", "Lucro Líquido", fill_lucro_liq),
        ("O1", "% líquida", fill_pct_liq),
    ]
    for ref, label, fill in headers_l1:
        if ":" in ref:
            ws.merge_cells(ref)
            anchor = ref.split(":")[0]
        else:
            anchor = ref
        c = ws[anchor]
        c.value = label
        c.font = font_header
        c.alignment = center
        c.fill = fill

    # === Linha 2: subcabeçalhos ===
    headers_l2 = {
        "A2": "SKU",
        "B2": "ID do Produto",
        "C2": "Valor de venda",
        "D2": "Clássico",
        "E2": "Premium",
        "L2": "%",
        "M2": "R$",
    }
    for cell, value in headers_l2.items():
        ws[cell] = value
        ws[cell].font = font_header
        ws[cell].alignment = center
        ws[cell].fill = fill_header_yellow

    # Bordas nas duas linhas de cabeçalho
    for row in (1, 2):
        for col in range(1, 16):  # A até O
            ws.cell(row=row, column=col).border = border

    # === Diagnóstico ===
    diagnostico: dict[str, Any] = {
        "matches": {
            "manual": 0,
            "exato": 0,
            "case_insensitive": 0,
            "prefixo": 0,
            "prefixo_divergente": 0,
            "nao_encontrado": 0,
        },
        "anuncios_sem_custo": [],
        "anuncios_com_fallback": [],
    }

    # === Linhas de dados ===
    for idx, r in enumerate(resultados, start=3):
        if "erro" in r:
            ws.cell(row=idx, column=1, value=f"ERRO: {r['erro']}")
            ws.cell(row=idx, column=2, value=r.get("item_id"))
            continue

        sku = r.get("sku") or "(sem SKU)"
        mlb = r["item_id"]
        preco = r["preco"]
        com_val = r.get("comissao_valor", 0) or 0
        com_pct = r.get("comissao_percentual", 0) or 0
        tf = r.get("tarifa_fixa", 0) or 0
        frete = r.get("frete_vendedor", 0) or 0
        modalidade = r.get("modalidade", "")

        # A: SKU  # noqa: ERA001
        ws.cell(row=idx, column=1, value=sku).alignment = left
        # B: MLB  # noqa: ERA001
        ws.cell(row=idx, column=2, value=mlb).alignment = left
        # C: Valor de venda
        c = ws.cell(row=idx, column=3, value=preco)
        c.number_format = 'R$ #,##0.00'

        # D ou E: Tarifa Clássico/Premium
        if modalidade == "Premium":
            ws.cell(row=idx, column=5, value=com_val).number_format = 'R$ #,##0.00'
        else:
            ws.cell(row=idx, column=4, value=com_val).number_format = 'R$ #,##0.00'

        # F: % da Tarifa
        ws.cell(row=idx, column=6, value=com_pct / 100 if com_pct else 0).number_format = '0.0%'
        # G: Tarifa Fixa
        ws.cell(row=idx, column=7, value=tf).number_format = 'R$ #,##0.00'
        # H: Frete  # noqa: ERA001
        ws.cell(row=idx, column=8, value=frete).number_format = 'R$ #,##0.00'
        # I: Valor Final (fórmula)
        ws.cell(
            row=idx, column=9,
            value=f"=C{idx}-D{idx}-E{idx}-G{idx}-H{idx}",
        ).number_format = 'R$ #,##0.00'

        # J: Custo do Produto (lookup com cascata)
        custo, fonte = buscar_custo(r.get("sku"), mlb, custos, custos_manuais)
        c_custo = ws.cell(row=idx, column=10)
        if custo is not None:
            c_custo.value = custo
            c_custo.number_format = 'R$ #,##0.00'
            if fonte == "manual":
                diagnostico["matches"]["manual"] += 1
            elif fonte == "exato":
                diagnostico["matches"]["exato"] += 1
            elif fonte.startswith("case_insensitive"):
                diagnostico["matches"]["case_insensitive"] += 1
                diagnostico["anuncios_com_fallback"].append(
                    {"item_id": mlb, "sku": sku, "fonte": fonte}
                )
            elif fonte.startswith("prefixo_divergente"):
                diagnostico["matches"]["prefixo_divergente"] += 1
                c_custo.fill = fill_warning
                diagnostico["anuncios_com_fallback"].append(
                    {"item_id": mlb, "sku": sku, "fonte": fonte}
                )
            elif fonte.startswith("prefixo"):
                diagnostico["matches"]["prefixo"] += 1
                diagnostico["anuncios_com_fallback"].append(
                    {"item_id": mlb, "sku": sku, "fonte": fonte}
                )
            # K: Lucro Bruto = Valor Final - Custo
            ws.cell(
                row=idx, column=11, value=f"=I{idx}-J{idx}"
            ).number_format = 'R$ #,##0.00'
            # N: Lucro Líquido = Lucro Bruto - Imposto R$
            ws.cell(
                row=idx, column=14, value=f"=K{idx}-M{idx}"
            ).number_format = 'R$ #,##0.00'
            # O: % líquida = Lucro Líquido / Valor de venda
            ws.cell(
                row=idx, column=15, value=f"=N{idx}/C{idx}"
            ).number_format = '0.00%'
        else:
            c_custo.fill = fill_warning
            c_custo.value = "(sem custo)"
            diagnostico["matches"]["nao_encontrado"] += 1
            diagnostico["anuncios_sem_custo"].append(
                {"item_id": mlb, "sku": sku, "titulo": r.get("title")}
            )

        # L: Imposto % (referência à célula R1)
        ws.cell(row=idx, column=12, value="=$R$1").number_format = '0.00%'
        # M: Imposto R$ = Valor de venda x Imposto %
        ws.cell(
            row=idx, column=13, value=f"=C{idx}*L{idx}"
        ).number_format = 'R$ #,##0.00'

        # Bordas e fonte na linha inteira
        for col in range(1, 16):
            cc = ws.cell(row=idx, column=col)
            cc.font = font_data
            cc.border = border

    # === Célula de configuração da alíquota (Q1/R1) ===
    ws["Q1"] = "Alíquota Imposto:"
    ws["Q1"].font = font_header
    ws["Q1"].alignment = right
    ws["R1"] = aliquota
    ws["R1"].number_format = '0.00%'
    ws["R1"].font = font_header
    ws["R1"].fill = fill_config
    ws["R1"].border = border
    ws["R1"].alignment = center

    # === Larguras de coluna ===
    larguras = {
        "A": 22, "B": 16, "C": 14, "D": 12, "E": 12,
        "F": 11, "G": 11, "H": 10, "I": 14,
        "J": 16, "K": 14, "L": 9, "M": 12, "N": 14, "O": 11,
        "Q": 18, "R": 12,
    }
    for col_letter, w in larguras.items():
        ws.column_dimensions[col_letter].width = w

    ws.row_dimensions[1].height = 24
    ws.row_dimensions[2].height = 20

    # Painéis congelados
    ws.freeze_panes = "A3"

    # Serializa pra bytes
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue(), diagnostico
