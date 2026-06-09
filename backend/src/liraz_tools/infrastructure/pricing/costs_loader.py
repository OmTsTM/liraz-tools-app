"""Carrega o mapa SKU → custo da planilha de custos.

Portado do MCP (`tools/relatorio.py` funções `_carregar_custos` e `_buscar_custo`),
com o termo "rinaldo" substituído por "custos" no app desktop. A lógica de
parsing e cascata de match é idêntica.

Schema esperado do XLSX:
- Aba 'Produtos'
- Coluna D = CÓDIGO (SKU)
- Coluna H = SOMA (custo)

Hierarquia de match:
  1) custos_manuais por MLB  (anúncios sem SKU)
  2) custos_manuais por SKU
  3) match exato
  4) case-insensitive
  5) prefixo antes do '-' (pra famílias 5209-10, 5209-23)
"""
from __future__ import annotations

from pathlib import Path

import openpyxl


class CustosXLSXError(Exception):
    """Erro ao carregar a planilha de custos."""


def carregar_custos(custos_xlsx_path: str | Path) -> dict[str, float]:
    """Lê o XLSX e retorna o mapa SKU → custo da aba 'Produtos'."""
    path = Path(custos_xlsx_path)
    if not path.exists():
        raise CustosXLSXError(
            f"Planilha de custos não encontrada: {path}"
        )

    try:
        wb = openpyxl.load_workbook(path, data_only=True)
    except Exception as e:
        raise CustosXLSXError(
            f"Falha ao abrir {path}: {e}"
        ) from e

    if "Produtos" not in wb.sheetnames:
        raise CustosXLSXError(
            f"Aba 'Produtos' não encontrada em {path.name}. "
            f"Abas disponíveis: {wb.sheetnames}"
        )

    ws = wb["Produtos"]
    custos: dict[str, float] = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        sku = row[3] if len(row) > 3 else None    # col D
        soma = row[7] if len(row) > 7 else None   # col H
        if sku is None or soma is None:
            continue
        # SKUs numéricos puros (ex: 1438 sem aspas no Excel) chegam como
        # int/float. Converte pra string e remove trailing ".0" de floats
        # acidentais — bug rev15: antes só aceitava `isinstance(sku, str)`,
        # ignorando silenciosamente ~93 SKUs Conquista numéricos.
        if isinstance(sku, (int, float)):
            sku_str = str(int(sku)) if float(sku).is_integer() else str(sku)
        elif isinstance(sku, str):
            sku_str = sku
        else:
            continue  # tipos exóticos (datetime, etc): pula
        sku_str = sku_str.strip()
        if not sku_str:
            continue
        try:
            custos[sku_str] = float(soma)
        except (TypeError, ValueError):
            continue

    return custos


def buscar_custo(
    sku: str | None,
    item_id: str,
    custos: dict[str, float],
    custos_manuais: dict[str, float] | None = None,
) -> tuple[float | None, str]:
    """Resolve custo com cascata: manual → exato → case_insensitive → prefixo.

    Returns:
        Tuple (valor, fonte). Fonte indica origem do match:
          - 'manual'
          - 'exato'
          - 'case_insensitive:<sku_real>'
          - 'prefixo:<base>'
          - 'prefixo_divergente:<base>'  (família com valores diferentes)
          - 'nao_encontrado'
    """
    cm = custos_manuais or {}

    # 0) override manual por MLB (anúncios sem SKU)
    if item_id in cm:
        return cm[item_id], "manual"

    # 1) override manual por SKU
    if sku and sku in cm:
        return cm[sku], "manual"

    if not sku:
        return None, "nao_encontrado"

    # 2) match exato
    if sku in custos:
        return custos[sku], "exato"

    # 3) case-insensitive
    custos_lower = {k.lower(): (k, v) for k, v in custos.items()}
    if sku.lower() in custos_lower:
        original, valor = custos_lower[sku.lower()]
        return valor, f"case_insensitive:{original}"

    # 4) fallback por prefixo
    if "-" in sku:
        prefixo = sku.split("-")[0]
        candidatos = [
            (k, v) for k, v in custos.items()
            if k.startswith(prefixo + "-") or k == prefixo
        ]
        if candidatos:
            valores = {round(v, 2) for _, v in candidatos}
            if len(valores) == 1:
                return candidatos[0][1], f"prefixo:{prefixo}"
            # Família com custos divergentes: usa primeiro e sinaliza
            return candidatos[0][1], f"prefixo_divergente:{prefixo}"

    return None, "nao_encontrado"


def carregar_tarifas_ml(
    tarifas_xlsx_path: str | Path | None,
) -> dict[str, dict[str, float]]:
    """Lê a planilha de tarifas reais por anúncio (formato da extensão Chrome).

    Schema esperado:
    - Aba 'TarifasML' (case-sensitive) OU primeira aba do arquivo
    - Coluna A = MLB (MLBxxxxxxxxxx)
    - Coluna B = Frete (R$, decimal) — vendedor paga frete (frete grátis / ≥R$79)
    - Coluna C = Custo fixo (R$, decimal) — tarifa fixa (envio por conta do comprador)
    - Cabeçalho na linha 1 (ignorado).

    Regra (combinada com a planilha exportada pela extensão):
      - Frete > 0      → Custo fixo = 0
      - Custo fixo > 0 → Frete = 0

    Retorna {item_id: {"custo_fixo": float, "frete": float}}. Path None, arquivo
    ausente ou aba vazia retorna dict vazio (override é OPCIONAL — sem ele o
    calculator usa o teto teórico).
    """
    if not tarifas_xlsx_path:
        return {}
    path = Path(tarifas_xlsx_path)
    if not path.exists():
        return {}

    try:
        wb = openpyxl.load_workbook(path, data_only=True)
    except Exception:
        return {}

    # Aceita aba 'TarifasML' (nome canônico exportado pela extensão) OU primeira
    # aba do arquivo (se o user editou manualmente e renomeou).
    if "TarifasML" in wb.sheetnames:
        ws = wb["TarifasML"]
    elif wb.sheetnames:
        ws = wb[wb.sheetnames[0]]
    else:
        return {}

    out: dict[str, dict[str, float]] = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row:
            continue
        item_id = row[0] if len(row) > 0 else None
        frete = row[1] if len(row) > 1 else None       # coluna B
        custo_fixo = row[2] if len(row) > 2 else None  # coluna C
        if not isinstance(item_id, str) or not item_id.startswith("MLB"):
            continue
        try:
            cf = float(custo_fixo) if custo_fixo is not None else 0.0
            fr = float(frete) if frete is not None else 0.0
        except (TypeError, ValueError):
            continue
        out[item_id.strip()] = {"custo_fixo": cf, "frete": fr}

    return out
